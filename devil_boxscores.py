"""Box scores for every Euroleague devil-week game since 2016-17.

Counting stats come from the official game-stats endpoint. Rates are computed
with the same possession, ORtg, DRtg, eFG% and TS% definitions as advanced.py.
Scheduled games stay blank: a missing box score is not filled in.
"""

from __future__ import annotations

import json
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from advanced import defensive_rating, effective_fg, offensive_rating, possessions, true_shooting

SOURCE = Path(__file__).with_name("euroleague-devil-weeks.xlsx")
CACHE = Path("/tmp/devil-boxscores")
OUTPUT = Path("/workspace/euroleague-devil-weeks-boxscores.xlsx")
STATS_URL = (
    "https://api-live.euroleague.net/v2/competitions/"
    "{competition}/seasons/{season_code}/games/{game_code}/stats"
)
USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)

COUNTING = [
    "PTS",
    "FGM",
    "FGA",
    "FG%",
    "3PM",
    "3PA",
    "3P%",
    "FTM",
    "FTA",
    "FT%",
    "ORB",
    "DRB",
    "TRB",
    "AST",
    "STL",
    "BLK",
    "TOV",
    "PF",
]
RATES = [
    "TS%",
    "eFG%",
    "ORB%",
    "DRB%",
    "TRB%",
    "AST%",
    "TOV%",
    "STL%",
    "BLK%",
    "PPS",
    "ORtg",
    "DRtg",
    "Pace",
]
IDENTITY = [
    "Σεζόν",
    "Γύρος",
    "Ημερομηνία",
    "Εβδομάδα",
    "Κωδικός αγώνα",
    "Έδρα",
    "Ομάδα",
    "Αντίπαλος",
    "Κατάσταση",
]
PERCENT = {"FG%", "3P%", "FT%", "TS%", "eFG%", "ORB%", "DRB%", "TRB%", "AST%", "TOV%", "STL%", "BLK%"}
ONE_DECIMAL = {"PPS", "ORtg", "DRtg", "Pace"}


def _ratio(numerator: float | None, denominator: float | None) -> float | None:
    if numerator is None or denominator is None or denominator == 0:
        return None
    return numerator / denominator


def _number(total: dict, key: str) -> float:
    value = total.get(key)
    if value is None:
        return 0.0
    return float(value)


def fetch_stats(game_code: str) -> dict | None:
    """Return the cached or freshly downloaded stats payload for one game code."""
    CACHE.mkdir(parents=True, exist_ok=True)
    path = CACHE / f"{game_code}.json"
    if path.exists():
        cached = json.loads(path.read_text())
        if cached.get("_missing"):
            return None
        return cached
    competition, season_code, number = game_code.split("_")[0][0], game_code.split("_")[0], game_code.split("_")[1]
    url = STATS_URL.format(competition=competition, season_code=season_code, game_code=number)
    last_error = ""
    for attempt in range(5):
        request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept": "application/json"})
        try:
            with urllib.request.urlopen(request, timeout=40) as response:
                payload = json.load(response)
            path.write_text(json.dumps(payload))
            return payload
        except urllib.error.HTTPError as exc:
            last_error = f"HTTP {exc.code}"
            if exc.code in {404, 400}:
                break
            if exc.code == 429:
                time.sleep(2.5 * (attempt + 1))
                continue
        except Exception as exc:  # noqa: BLE001
            last_error = str(exc)
        time.sleep(0.6 * (attempt + 1))
    if last_error.startswith("HTTP 429"):
        return None
    path.write_text(json.dumps({"_missing": True, "error": last_error}))
    return None


def side_line(total: dict, opponent: dict) -> dict:
    """Traditional totals plus the requested rates for one team in one game."""
    pts = _number(total, "points")
    fgm = _number(total, "fieldGoalsMadeTotal")
    fga = _number(total, "fieldGoalsAttemptedTotal")
    tpm = _number(total, "fieldGoalsMade3")
    tpa = _number(total, "fieldGoalsAttempted3")
    ftm = _number(total, "freeThrowsMade")
    fta = _number(total, "freeThrowsAttempted")
    orb = _number(total, "offensiveRebounds")
    drb = _number(total, "defensiveRebounds")
    trb = _number(total, "totalRebounds")
    ast = _number(total, "assistances")
    stl = _number(total, "steals")
    blk = _number(total, "blocksFavour")
    tov = _number(total, "turnovers")
    pf = _number(total, "foulsCommited")
    opp_pts = _number(opponent, "points")
    opp_orb = _number(opponent, "offensiveRebounds")
    opp_drb = _number(opponent, "defensiveRebounds")
    opp_trb = _number(opponent, "totalRebounds")
    opp_fga = _number(opponent, "fieldGoalsAttemptedTotal")
    opp_tpa = _number(opponent, "fieldGoalsAttempted3")
    opp_two_attempts = opponent.get("fieldGoalsAttempted2")
    opp_fga2 = float(opp_two_attempts) if opp_two_attempts is not None else opp_fga - opp_tpa
    poss = possessions(fga, orb, tov, fta)
    opp_poss = possessions(
        opp_fga,
        opp_orb,
        _number(opponent, "turnovers"),
        _number(opponent, "freeThrowsAttempted"),
    )
    game_minutes = _number(total, "timePlayed") / 60.0
    pace = None
    if poss is not None and opp_poss is not None and game_minutes:
        pace = 40.0 * ((poss + opp_poss) / 2.0) / game_minutes
    return {
        "PTS": pts,
        "FGM": fgm,
        "FGA": fga,
        "FG%": _ratio(fgm, fga),
        "3PM": tpm,
        "3PA": tpa,
        "3P%": _ratio(tpm, tpa),
        "FTM": ftm,
        "FTA": fta,
        "FT%": _ratio(ftm, fta),
        "ORB": orb,
        "DRB": drb,
        "TRB": trb,
        "AST": ast,
        "STL": stl,
        "BLK": blk,
        "TOV": tov,
        "PF": pf,
        "TS%": true_shooting(pts, fga, fta),
        "eFG%": effective_fg(fgm, tpm, fga),
        "ORB%": _ratio(orb, orb + opp_drb),
        "DRB%": _ratio(drb, drb + opp_orb),
        "TRB%": _ratio(trb, trb + opp_trb),
        "AST%": _ratio(ast, fgm),
        "TOV%": _ratio(tov, fga + 0.44 * fta + tov),
        "STL%": _ratio(stl, opp_poss),
        "BLK%": _ratio(blk, opp_fga2),
        "PPS": _ratio(pts, fga),
        "ORtg": offensive_rating(pts, poss),
        "DRtg": defensive_rating(opp_pts, poss),
        "Pace": pace,
    }


def empty_line() -> dict:
    return {column: None for column in COUNTING + RATES}


def build() -> None:
    games = pd.read_excel(SOURCE)
    played = games[games["Κατάσταση"] == "Παίχτηκε"]
    codes = played["Κωδικός αγώνα"].astype(str).tolist()
    print(f"fetching {len(codes)} played games", flush=True)
    payloads: dict[str, dict | None] = {}
    with ThreadPoolExecutor(max_workers=3) as pool:
        futures = {pool.submit(fetch_stats, code): code for code in codes}
        done = 0
        for future in as_completed(futures):
            code = futures[future]
            payloads[code] = future.result()
            done += 1
            if done % 100 == 0 or done == len(codes):
                missing = sum(1 for value in payloads.values() if value is None)
                print(f"{done}/{len(codes)} missing {missing}", flush=True)

    rows: list[dict] = []
    mismatches = 0
    missing_played = 0
    for record in games.to_dict("records"):
        code = str(record["Κωδικός αγώνα"])
        payload = payloads.get(code)
        home_name = record["Γηπεδούχος"]
        away_name = record["Φιλοξενούμενη"]
        identity = {
            "Σεζόν": record["Σεζόν"],
            "Γύρος": record["Γύρος"],
            "Ημερομηνία": record["Ημερομηνία"],
            "Εβδομάδα": record["Εβδομάδα"],
            "Κωδικός αγώνα": code,
            "Κατάσταση": record["Κατάσταση"],
        }
        if payload and payload.get("local") and payload.get("road"):
            local = (payload.get("local") or {}).get("total") or {}
            road = (payload.get("road") or {}).get("total") or {}
            if not local or not road:
                missing_played += 1
                lines = (empty_line(), empty_line())
            else:
                workbook_home = record["Πόντοι γηπεδούχου"]
                workbook_away = record["Πόντοι φιλοξενούμενης"]
                if workbook_home is not None and float(local.get("points") or 0) != float(workbook_home):
                    mismatches += 1
                if workbook_away is not None and float(road.get("points") or 0) != float(workbook_away):
                    mismatches += 1
                lines = (side_line(local, road), side_line(road, local))
        else:
            if record["Κατάσταση"] == "Παίχτηκε":
                missing_played += 1
            lines = (empty_line(), empty_line())
        for venue, team, opponent, line in (
            ("Γηπεδούχος", home_name, away_name, lines[0]),
            ("Φιλοξενούμενη", away_name, home_name, lines[1]),
        ):
            rows.append({**identity, "Έδρα": venue, "Ομάδα": team, "Αντίπαλος": opponent, **line})

    frame = pd.DataFrame(rows, columns=IDENTITY + COUNTING + RATES)
    notes = pd.DataFrame(
        [
            ["Αγώνες στο αρχείο διαβολοβδομάδων", len(games)],
            ["Αγώνες με box score", int((frame["PTS"].notna() & (frame["Έδρα"] == "Γηπεδούχος")).sum())],
            ["Αγώνες χωρίς box score", int((frame["PTS"].isna() & (frame["Έδρα"] == "Γηπεδούχος")).sum())],
            ["Από αυτούς, προγραμματισμένοι", int((games["Κατάσταση"] != "Παίχτηκε").sum())],
            ["Παιγμένοι χωρίς απάντηση API", missing_played],
            ["Γραμμές με απόκλιση πόντων από το αρχείο αγώνων", mismatches],
            ["Πηγή", "https://api-live.euroleague.net/v2/competitions/E/seasons/{σεζόν}/games/{κωδικός}/stats"],
            ["Κατοχή", "0,96 × (FGA − ORB + TOV + 0,44 × FTA)"],
            ["FG%", "FGM / FGA"],
            ["3P%", "3PM / 3PA"],
            ["FT%", "FTM / FTA"],
            ["TS%", "PTS / (2 × (FGA + 0,44 × FTA))"],
            ["eFG%", "(FGM + 0,5 × 3PM) / FGA"],
            ["ORB%", "ORB / (ORB + DRB αντιπάλου)"],
            ["DRB%", "DRB / (DRB + ORB αντιπάλου)"],
            ["TRB%", "TRB / (TRB + TRB αντιπάλου)"],
            ["AST%", "AST / FGM"],
            ["TOV%", "TOV / (FGA + 0,44 × FTA + TOV)"],
            ["STL%", "STL / κατοχή αντιπάλου"],
            ["BLK%", "BLK / δίποντα που επιχείρησε ο αντίπαλος"],
            ["PPS", "PTS / FGA"],
            ["ORtg", "PTS × 100 / κατοχή της ομάδας"],
            ["DRtg", "PTS αντιπάλου × 100 / κατοχή της ομάδας"],
            ["Pace", "40 × ((κατοχή + κατοχή αντιπάλου) / 2) / λεπτά αγώνα"],
            ["Κενά κελιά", "Προγραμματισμένοι αγώνες και αγώνες χωρίς box score. Δεν συμπληρώθηκαν αριθμοί."],
        ],
        columns=["Πεδίο", "Τιμή"],
    )
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Στατιστικά", index=False)
        notes.to_excel(writer, sheet_name="Σημείωση", index=False)
        sheet = writer.sheets["Στατιστικά"]
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 30
        widths = {
            "Σεζόν": 12,
            "Γύρος": 14,
            "Ημερομηνία": 14,
            "Εβδομάδα": 12,
            "Κωδικός αγώνα": 16,
            "Έδρα": 16,
            "Ομάδα": 32,
            "Αντίπαλος": 32,
            "Κατάσταση": 20,
        }
        for index, name in enumerate(frame.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 10)
            for row in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell = row[0]
                if name == "Ημερομηνία" and cell.value is not None:
                    cell.number_format = "YYYY-MM-DD"
                elif name in PERCENT and cell.value is not None:
                    cell.number_format = "0.0%"
                elif name in ONE_DECIMAL and cell.value is not None:
                    cell.number_format = "0.0"
                elif name in COUNTING and cell.value is not None:
                    cell.number_format = "0"
        note_sheet = writer.sheets["Σημείωση"]
        note_sheet.column_dimensions["A"].width = 52
        note_sheet.column_dimensions["B"].width = 88
        for cell in note_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
    print(
        f"wrote {OUTPUT} rows {len(frame)} missing_played {missing_played} mismatches {mismatches}",
        flush=True,
    )


if __name__ == "__main__":
    build()
