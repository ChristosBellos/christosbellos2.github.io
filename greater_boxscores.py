"""Box scores for the Euroleague and Eurocup games listed in greater_games.xlsx.

Home and away are separate rows. Counting stats come from the official game
feed. Rates use the same formulas as devil_boxscores.py. A game with no box
score keeps its identity and blank stat cells.
"""

from __future__ import annotations

import json
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from devil_boxscores import (
    COUNTING,
    ONE_DECIMAL,
    PERCENT,
    RATES,
    USER_AGENT,
    empty_line,
    side_line,
)

SOURCE = Path(__file__).with_name("greater_games.xlsx")
OUTPUT = Path(__file__).with_name("greater_games-boxscores.xlsx")
CACHE = Path("/tmp/greater-boxscores")
EXTRA_CACHE = Path("/tmp/devil-boxscores")
STATS_URL = (
    "https://api-live.euroleague.net/v2/competitions/"
    "{competition}/seasons/{season_code}/games/{game_code}/stats"
)
GAMES_URL = (
    "https://api-live.euroleague.net/v2/competitions/"
    "{competition}/seasons/{season_code}/games?limit=600"
)

IDENTITY = [
    "Διοργάνωση",
    "Σεζόν",
    "Έτος λήξης",
    "Φάση",
    "Γύρος",
    "Όμιλος",
    "Ημερομηνία",
    "Κωδικός αγώνα",
    "Mean Day Difference",
    "Day Difference",
    "Έδρα",
    "Ομάδα",
    "Αντίπαλος",
]


def _get(url: str) -> dict:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept": "application/json"})
    with urllib.request.urlopen(request, timeout=40) as response:
        return json.load(response)


def fetch_stats(game_code: str) -> dict | None:
    CACHE.mkdir(parents=True, exist_ok=True)
    path = CACHE / f"{game_code}.json"
    for candidate in (path, EXTRA_CACHE / f"{game_code}.json"):
        if not candidate.exists():
            continue
        cached = json.loads(candidate.read_text())
        if cached.get("_missing"):
            continue
        if candidate != path:
            path.write_text(json.dumps(cached))
        return cached
    competition, season_code, number = game_code[0], game_code.split("_")[0], game_code.split("_")[1]
    url = STATS_URL.format(competition=competition, season_code=season_code, game_code=number)
    last_error = ""
    for attempt in range(5):
        try:
            payload = _get(url)
            path.write_text(json.dumps(payload))
            return payload
        except urllib.error.HTTPError as exc:
            last_error = f"HTTP {exc.code}"
            if exc.code in {404, 400}:
                break
            if exc.code == 429:
                time.sleep(2.5 * (attempt + 1))
                continue
        except Exception as exc:  # noqa: BLE001
            last_error = str(exc)
        time.sleep(0.6 * (attempt + 1))
    if last_error.startswith("HTTP 429"):
        return None
    path.write_text(json.dumps({"_missing": True, "error": last_error}))
    return None


def _joined(values: pd.Series) -> object:
    kept: list[str] = []
    for value in values:
        if pd.isna(value):
            continue
        text = str(value).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            text = text[:-2]
        if text and text not in kept:
            kept.append(text)
    if not kept:
        return None
    if len(kept) == 1:
        try:
            return float(kept[0]) if "." in kept[0] else int(kept[0])
        except ValueError:
            return kept[0]
    return " | ".join(kept)


def game_table() -> tuple[pd.DataFrame, int]:
    raw = pd.read_excel(SOURCE)
    raw["Κωδικός αγώνα"] = raw["Κωδικός αγώνα"].astype(str)
    rows = []
    for code, group in raw.groupby("Κωδικός αγώνα", sort=False):
        first = group.iloc[0]
        rows.append(
            {
                "Διοργάνωση": first["Organization"],
                "Σεζόν": first["Season"],
                "Έτος λήξης": int(first["End Year"]),
                "Φάση": first["Phase_x"],
                "Γύρος": first["Round"],
                "Όμιλος": first["Group"],
                "Ημερομηνία": first["Date"],
                "Κωδικός αγώνα": code,
                "Mean Day Difference": _joined(group["Mean Day Difference"]),
                "Day Difference": _joined(group["Day Difference"]),
            }
        )
    return pd.DataFrame(rows), len(raw)


def club_names(frame: pd.DataFrame) -> dict[str, tuple[str, str]]:
    seasons = sorted({code.split("_")[0] for code in frame["Κωδικός αγώνα"]})
    found: dict[str, tuple[str, str]] = {}
    for season_code in seasons:
        competition = season_code[0]
        try:
            payload = _get(GAMES_URL.format(competition=competition, season_code=season_code))
        except Exception as exc:  # noqa: BLE001
            print(f"games list {season_code} failed: {exc}", flush=True)
            continue
        for game in payload.get("data") or []:
            code = f"{season_code}_{game.get('gameCode')}"
            local = ((game.get("local") or {}).get("club") or {}).get("name") or ""
            road = ((game.get("road") or {}).get("club") or {}).get("name") or ""
            if local or road:
                found[code] = (str(local), str(road))
        print(f"names {season_code}: {len(payload.get('data') or [])}", flush=True)
    return found


def has_box_score(total: dict | None) -> bool:
    """A walkover can return an all-zero total and no player lines."""
    if not total:
        return False
    counted = (
        "points",
        "fieldGoalsAttemptedTotal",
        "freeThrowsAttempted",
        "totalRebounds",
        "turnovers",
    )
    return any(total.get(key) not in (None, 0, 0.0) for key in counted)


def load_stats(codes: list[str]) -> dict[str, dict | None]:
    payloads: dict[str, dict | None] = {}
    print(f"fetching {len(codes)} box scores", flush=True)
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures = {pool.submit(fetch_stats, code): code for code in codes}
        done = 0
        for future in as_completed(futures):
            payloads[futures[future]] = future.result()
            done += 1
            if done % 100 == 0 or done == len(codes):
                missing = sum(value is None for value in payloads.values())
                print(f"{done}/{len(codes)} pending {missing}", flush=True)
    return payloads


def build() -> None:
    games, source_rows = game_table()
    names = club_names(games)
    codes = games["Κωδικός αγώνα"].tolist()
    payloads = load_stats(codes)
    pending = [code for code in codes if payloads.get(code) is None and not (CACHE / f"{code}.json").exists()]
    if pending:
        print(f"retrying {len(pending)}", flush=True)
        time.sleep(8)
        payloads.update(load_stats(pending))

    rows: list[dict] = []
    missing = 0
    for record in games.to_dict("records"):
        code = record["Κωδικός αγώνα"]
        home_name, away_name = names.get(code, ("", ""))
        payload = payloads.get(code)
        local = ((payload or {}).get("local") or {}).get("total") or {}
        road = ((payload or {}).get("road") or {}).get("total") or {}
        if has_box_score(local) and has_box_score(road):
            lines = (side_line(local, road), side_line(road, local))
        else:
            missing += 1
            lines = (empty_line(), empty_line())
        for venue, team, opponent, line in (
            ("Γηπεδούχος", home_name, away_name, lines[0]),
            ("Φιλοξενούμενη", away_name, home_name, lines[1]),
        ):
            rows.append({**record, "Έδρα": venue, "Ομάδα": team, "Αντίπαλος": opponent, **line})

    frame = pd.DataFrame(rows, columns=IDENTITY + COUNTING + RATES)
    notes = pd.DataFrame(
        [
            ["Γραμμές στο αρχείο που δόθηκε", source_rows],
            ["Μοναδικοί αγώνες", len(games)],
            ["Αγώνες με box score", int((frame["PTS"].notna() & (frame["Έδρα"] == "Γηπεδούχος")).sum())],
            ["Αγώνες χωρίς box score", missing],
            ["Γραμμές στο αποτέλεσμα", len(frame)],
            ["Διπλές γραμμές στο αρχείο εισόδου", "Ο ίδιος κωδικός αγώνα επαναλαμβάνεται. Μία φορά στο αποτέλεσμα. Αν το Day Difference διέφερε, οι τιμές ενώθηκαν με ' | '."],
            ["Σειρά γραμμών", "Πρώτα η γηπεδούχος, αμέσως μετά η φιλοξενούμενη."],
            ["Πηγή", "https://api-live.euroleague.net/v2/competitions/{E ή U}/seasons/{κωδικός}/games/{αγώνας}/stats"],
            ["Κατοχή", "0,96 × (FGA − ORB + TOV + 0,44 × FTA)"],
            ["FG%", "FGM / FGA"],
            ["3P%", "3PM / 3PA"],
            ["FT%", "FTM / FTA"],
            ["TS%", "PTS / (2 × (FGA + 0,44 × FTA))"],
            ["eFG%", "(FGM + 0,5 × 3PM) / FGA"],
            ["ORB%", "ORB / (ORB + DRB αντιπάλου)"],
            ["DRB%", "DRB / (DRB + ORB αντιπάλου)"],
            ["TRB%", "TRB / (TRB + TRB αντιπάλου)"],
            ["AST%", "AST / FGM"],
            ["TOV%", "TOV / (FGA + 0,44 × FTA + TOV)"],
            ["STL%", "STL / κατοχή αντιπάλου"],
            ["BLK%", "BLK / δίποντα που επιχείρησε ο αντίπαλος"],
            ["PPS", "PTS / FGA"],
            ["ORtg", "PTS × 100 / κατοχή της ομάδας"],
            ["DRtg", "PTS αντιπάλου × 100 / κατοχή της ομάδας"],
            ["Pace", "40 × ((κατοχή + κατοχή αντιπάλου) / 2) / λεπτά αγώνα"],
            ["Κενά κελιά", "Αγώνες χωρίς box score, ανάμεσά τους ακυρωμένοι και walkover χωρίς στατιστικά παικτών. Δεν συμπληρώθηκαν αριθμοί."],
        ],
        columns=["Πεδίο", "Τιμή"],
    )
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Στατιστικά", index=False)
        notes.to_excel(writer, sheet_name="Σημείωση", index=False)
        sheet = writer.sheets["Στατιστικά"]
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 30
        widths = {
            "Διοργάνωση": 14,
            "Σεζόν": 12,
            "Έτος λήξης": 12,
            "Φάση": 16,
            "Γύρος": 14,
            "Όμιλος": 18,
            "Ημερομηνία": 14,
            "Κωδικός αγώνα": 16,
            "Mean Day Difference": 22,
            "Day Difference": 18,
            "Έδρα": 16,
            "Ομάδα": 32,
            "Αντίπαλος": 32,
        }
        for index, name in enumerate(frame.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 10)
            for row in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell = row[0]
                if name == "Ημερομηνία" and cell.value is not None:
                    cell.number_format = "YYYY-MM-DD"
                elif name in PERCENT and isinstance(cell.value, float):
                    cell.number_format = "0.0%"
                elif name in ONE_DECIMAL and isinstance(cell.value, float):
                    cell.number_format = "0.0"
                elif name in COUNTING and isinstance(cell.value, float):
                    cell.number_format = "0"
        note_sheet = writer.sheets["Σημείωση"]
        note_sheet.column_dimensions["A"].width = 42
        note_sheet.column_dimensions["B"].width = 110
        for cell in note_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
    print(f"wrote {OUTPUT} rows {len(frame)} missing {missing}", flush=True)


if __name__ == "__main__":
    build()
