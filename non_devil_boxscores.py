"""Box scores for Euroleague games since 2016-17 that are not devil-week games.

The devil-week list is the uploaded workbook. Every other Euroleague game in
seasons E2016 through E2026 is included, regular season and the later rounds.
Home and away are separate rows. Games without a box score stay blank.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from devil_boxscores import COUNTING, ONE_DECIMAL, PERCENT, RATES, empty_line, side_line
from greater_boxscores import CACHE, _get, has_box_score, load_stats

DEVIL = Path(__file__).with_name("euroleague-devil-weeks.xlsx")
OUTPUT = Path(__file__).with_name("euroleague-non-devil-boxscores.xlsx")
GAMES_URL = (
    "https://api-live.euroleague.net/v2/competitions/E/seasons/{season_code}/games?limit=600"
)
PHASES = {
    "RS": "Regular Season",
    "PO": "Playoffs",
    "FF": "Final Four",
    "PI": "Play-In",
}

IDENTITY = [
    "Σεζόν",
    "Έτος λήξης",
    "Φάση",
    "Γύρος",
    "Ημερομηνία",
    "Ώρα",
    "Κωδικός αγώνα",
    "Κατάσταση",
    "Έδρα",
    "Ομάδα",
    "Αντίπαλος",
]


def season_label(season_code: str) -> tuple[str, int]:
    start = int(season_code[1:])
    return f"{start}-{str(start + 1)[-2:]}", start + 1


def devil_codes() -> set[str]:
    frame = pd.read_excel(DEVIL)
    return set(frame["Κωδικός αγώνα"].astype(str))


def calendar(excluded: set[str]) -> pd.DataFrame:
    rows = []
    for start in range(2016, 2027):
        season_code = f"E{start}"
        payload = _get(GAMES_URL.format(season_code=season_code))
        label, end_year = season_label(season_code)
        for game in payload.get("data") or []:
            code = f"{season_code}_{game.get('gameCode')}"
            if code in excluded:
                continue
            phase = (game.get("phaseType") or {}).get("code") or ""
            local = game.get("local") or {}
            road = game.get("road") or {}
            local_date = str(game.get("localDate") or game.get("utcDate") or "")
            status = "Παίχτηκε" if game.get("played") else "Προγραμματισμένος"
            game_status = str(game.get("gameStatus") or "")
            if game_status.lower() in {"cancelled", "canceled", "aborted"}:
                status = "Ακυρώθηκε"
            rows.append(
                {
                    "Σεζόν": label,
                    "Έτος λήξης": end_year,
                    "Φάση": PHASES.get(phase, (game.get("phaseType") or {}).get("name") or phase),
                    "Γύρος": game.get("roundName") or "",
                    "Ημερομηνία": local_date[:10],
                    "Ώρα": game.get("confirmedHour") or (local_date[11:16] if len(local_date) >= 16 else None),
                    "Κωδικός αγώνα": code,
                    "Κατάσταση": status,
                    "Γηπεδούχος": ((local.get("club") or {}).get("name") or ""),
                    "Φιλοξενούμενη": ((road.get("club") or {}).get("name") or ""),
                    "played": bool(game.get("played")),
                }
            )
        print(f"{season_code} kept {sum(row['Σεζόν']==label for row in rows)}", flush=True)
    frame = pd.DataFrame(rows)
    return frame.sort_values(["Έτος λήξης", "Ημερομηνία", "Κωδικός αγώνα"], kind="stable")


def build() -> None:
    excluded = devil_codes()
    games = calendar(excluded)
    played_codes = games.loc[games["played"], "Κωδικός αγώνα"].tolist()
    payloads = load_stats(played_codes)
    pending = [code for code in played_codes if payloads.get(code) is None and not (CACHE / f"{code}.json").exists()]
    if pending:
        print(f"retrying {len(pending)}", flush=True)
        payloads.update(load_stats(pending))

    rows: list[dict] = []
    missing = 0
    for record in games.to_dict("records"):
        code = record["Κωδικός αγώνα"]
        home_name = record.pop("Γηπεδούχος")
        away_name = record.pop("Φιλοξενούμενη")
        record.pop("played")
        payload = payloads.get(code)
        local = ((payload or {}).get("local") or {}).get("total") or {}
        road = ((payload or {}).get("road") or {}).get("total") or {}
        if has_box_score(local) and has_box_score(road):
            lines = (side_line(local, road), side_line(road, local))
        else:
            missing += 1
            lines = (empty_line(), empty_line())
        for venue, team, opponent, line in (
            ("Γηπεδούχος", home_name, away_name, lines[0]),
            ("Φιλοξενούμενη", away_name, home_name, lines[1]),
        ):
            rows.append({**record, "Έδρα": venue, "Ομάδα": team, "Αντίπαλος": opponent, **line})

    frame = pd.DataFrame(rows, columns=IDENTITY + COUNTING + RATES)
    notes = pd.DataFrame(
        [
            ["Ορισμός", "Αγώνες Euroleague από το 2016-17 που δεν είναι στον κατάλογο διαβολοβδομάδων. Η κανονική περίοδος, τα play-in, τα playoffs και το Final Four περιλαμβάνονται."],
            ["Αγώνες διαβολοβδομάδας που εξαιρέθηκαν", len(excluded)],
            ["Αγώνες στο αποτέλεσμα", len(games)],
            ["Αγώνες με box score", int((frame["PTS"].notna() & (frame["Έδρα"] == "Γηπεδούχος")).sum())],
            ["Αγώνες χωρίς box score", missing],
            ["Γραμμές", len(frame)],
            ["Σειρά γραμμών", "Πρώτα η γηπεδούχος, αμέσως μετά η φιλοξενούμενη."],
            ["Πηγή", "https://api-live.euroleague.net/v2/competitions/E/seasons/E2016/games/{αγώνας}/stats"],
            ["Κατοχή", "0,96 × (FGA − ORB + TOV + 0,44 × FTA)"],
            ["FG% / 3P% / FT%", "Εύστοχα / προσπάθειες"],
            ["TS%", "PTS / (2 × (FGA + 0,44 × FTA))"],
            ["eFG%", "(FGM + 0,5 × 3PM) / FGA"],
            ["ORB%", "ORB / (ORB + DRB αντιπάλου)"],
            ["DRB%", "DRB / (DRB + ORB αντιπάλου)"],
            ["TRB%", "TRB / (TRB + TRB αντιπάλου)"],
            ["AST%", "AST / FGM"],
            ["TOV%", "TOV / (FGA + 0,44 × FTA + TOV)"],
            ["STL%", "STL / κατοχή αντιπάλου"],
            ["BLK%", "BLK / δίποντα που επιχείρησε ο αντίπαλος"],
            ["PPS", "PTS / FGA"],
            ["ORtg", "PTS × 100 / κατοχή της ομάδας"],
            ["DRtg", "PTS αντιπάλου × 100 / κατοχή της ομάδας"],
            ["Pace", "40 × ((κατοχή + κατοχή αντιπάλου) / 2) / λεπτά αγώνα"],
            ["Κενά κελιά", "Προγραμματισμένοι, ακυρωμένοι και walkover χωρίς στατιστικά παικτών. Δεν συμπληρώθηκαν αριθμοί."],
        ],
        columns=["Πεδίο", "Τιμή"],
    )
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Στατιστικά", index=False)
        notes.to_excel(writer, sheet_name="Σημείωση", index=False)
        sheet = writer.sheets["Στατιστικά"]
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 30
        widths = {
            "Σεζόν": 12,
            "Έτος λήξης": 12,
            "Φάση": 16,
            "Γύρος": 22,
            "Ημερομηνία": 14,
            "Ώρα": 12,
            "Κωδικός αγώνα": 16,
            "Κατάσταση": 20,
            "Έδρα": 16,
            "Ομάδα": 32,
            "Αντίπαλος": 32,
        }
        for index, name in enumerate(frame.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 10)
            for row in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell = row[0]
                if name == "Ημερομηνία" and cell.value is not None:
                    cell.number_format = "YYYY-MM-DD"
                elif name in PERCENT and isinstance(cell.value, float):
                    cell.number_format = "0.0%"
                elif name in ONE_DECIMAL and isinstance(cell.value, float):
                    cell.number_format = "0.0"
                elif name in COUNTING and isinstance(cell.value, float):
                    cell.number_format = "0"
        note_sheet = writer.sheets["Σημείωση"]
        note_sheet.column_dimensions["A"].width = 48
        note_sheet.column_dimensions["B"].width = 120
        for cell in note_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
    print(f"wrote {OUTPUT} rows {len(frame)} missing {missing}", flush=True)


if __name__ == "__main__":
    build()
