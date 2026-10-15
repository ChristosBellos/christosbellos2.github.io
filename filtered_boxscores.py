"""Box scores for every Euroleague game listed in filtered_games.xlsx.

Home and away are separate rows. Counting stats come from the official game
feed. Rates use the same formulas as greater_boxscores.py. Games without a
box score keep their identity and blank stat cells.
"""

from __future__ import annotations

import time
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from devil_boxscores import COUNTING, ONE_DECIMAL, PERCENT, RATES, empty_line, side_line
from greater_boxscores import CACHE, club_names, has_box_score, load_stats

SOURCE = Path(__file__).with_name("filtered_games.xlsx")
OUTPUT = Path(__file__).with_name("filtered_games-boxscores.xlsx")

IDENTITY = [
    "Διοργάνωση",
    "Σεζόν",
    "Έτος λήξης",
    "Φάση",
    "Γύρος",
    "Όμιλος",
    "Ημερομηνία",
    "Ώρα",
    "Κωδικός αγώνα",
    "Mean Day Difference",
    "Day Difference",
    "Έδρα",
    "Ομάδα",
    "Αντίπαλος",
]


def game_table() -> pd.DataFrame:
    raw = pd.read_excel(SOURCE)
    raw["Κωδικός αγώνα"] = raw["Κωδικός αγώνα"].astype(str)
    rows = []
    for record in raw.to_dict("records"):
        rows.append(
            {
                "Διοργάνωση": record["Organization"],
                "Σεζόν": record["Season"],
                "Έτος λήξης": int(record["End Year"]),
                "Φάση": record["Phase_x"],
                "Γύρος": record["Round"],
                "Όμιλος": record["Group"],
                "Ημερομηνία": record["Date"],
                "Ώρα": record["Time"],
                "Κωδικός αγώνα": record["Κωδικός αγώνα"],
                "Mean Day Difference": record["Mean Day Difference"],
                "Day Difference": record["Day Difference"],
            }
        )
    return pd.DataFrame(rows)


def build() -> None:
    games = game_table()
    names = club_names(games)
    codes = games["Κωδικός αγώνα"].tolist()
    payloads = load_stats(codes)
    pending = [code for code in codes if payloads.get(code) is None and not (CACHE / f"{code}.json").exists()]
    if pending:
        print(f"retrying {len(pending)}", flush=True)
        time.sleep(8)
        payloads.update(load_stats(pending))
    rows: list[dict] = []
    missing = 0
    for record in games.to_dict("records"):
        code = record["Κωδικός αγώνα"]
        home_name, away_name = names.get(code, ("", ""))
        payload = payloads.get(code)
        local = ((payload or {}).get("local") or {}).get("total") or {}
        road = ((payload or {}).get("road") or {}).get("total") or {}
        if has_box_score(local) and has_box_score(road):
            lines = (side_line(local, road), side_line(road, local))
        else:
            missing += 1
            lines = (empty_line(), empty_line())
        for venue, team, opponent, line in (
            ("Γηπεδούχος", home_name, away_name, lines[0]),
            ("Φιλοξενούμενη", away_name, home_name, lines[1]),
        ):
            rows.append({**record, "Έδρα": venue, "Ομάδα": team, "Αντίπαλος": opponent, **line})

    frame = pd.DataFrame(rows, columns=IDENTITY + COUNTING + RATES)
    notes = pd.DataFrame(
        [
            ["Αγώνες στο αρχείο που δόθηκε", len(games)],
            ["Αγώνες με box score", int((frame["PTS"].notna() & (frame["Έδρα"] == "Γηπεδούχος")).sum())],
            ["Αγώνες χωρίς box score", missing],
            ["Γραμμές στο αποτέλεσμα", len(frame)],
            ["Σειρά γραμμών", "Πρώτα η γηπεδούχος, αμέσως μετά η φιλοξενούμενη."],
            ["Φάση", "Οι στήλες Phase_x και Phase_y του αρχείου εισόδου ήταν ίδιες, οπότε μένει μία στήλη."],
            ["Πηγή", "https://api-live.euroleague.net/v2/competitions/E/seasons/{κωδικός}/games/{αγώνας}/stats"],
            ["Κατοχή", "0,96 × (FGA − ORB + TOV + 0,44 × FTA)"],
            ["FG%", "FGM / FGA"],
            ["3P%", "3PM / 3PA"],
            ["FT%", "FTM / FTA"],
            ["TS%", "PTS / (2 × (FGA + 0,44 × FTA))"],
            ["eFG%", "(FGM + 0,5 × 3PM) / FGA"],
            ["ORB%", "ORB / (ORB + DRB αντιπάλου)"],
            ["DRB%", "DRB / (DRB + ORB αντιπάλου)"],
            ["TRB%", "TRB / (TRB + TRB αντιπάλου)"],
            ["AST%", "AST / FGM"],
            ["TOV%", "TOV / (FGA + 0,44 × FTA + TOV)"],
            ["STL%", "STL / κατοχή αντιπάλου"],
            ["BLK%", "BLK / δίποντα που επιχείρησε ο αντίπαλος"],
            ["PPS", "PTS / FGA"],
            ["ORtg", "PTS × 100 / κατοχή της ομάδας"],
            ["DRtg", "PTS αντιπάλου × 100 / κατοχή της ομάδας"],
            ["Pace", "40 × ((κατοχή + κατοχή αντιπάλου) / 2) / λεπτά αγώνα"],
            ["Κενά κελιά", "Αγώνες χωρίς box score, ανάμεσά τους ακυρωμένοι και walkover χωρίς στατιστικά παικτών. Δεν συμπληρώθηκαν αριθμοί."],
        ],
        columns=["Πεδίο", "Τιμή"],
    )
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Στατιστικά", index=False)
        notes.to_excel(writer, sheet_name="Σημείωση", index=False)
        sheet = writer.sheets["Στατιστικά"]
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 30
        widths = {
            "Διοργάνωση": 14,
            "Σεζόν": 12,
            "Έτος λήξης": 12,
            "Φάση": 16,
            "Γύρος": 14,
            "Όμιλος": 18,
            "Ημερομηνία": 14,
            "Ώρα": 12,
            "Κωδικός αγώνα": 16,
            "Mean Day Difference": 22,
            "Day Difference": 16,
            "Έδρα": 16,
            "Ομάδα": 32,
            "Αντίπαλος": 32,
        }
        for index, name in enumerate(frame.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 10)
            for row in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell = row[0]
                if name == "Ημερομηνία" and cell.value is not None:
                    cell.number_format = "YYYY-MM-DD"
                elif name in PERCENT and isinstance(cell.value, float):
                    cell.number_format = "0.0%"
                elif name in ONE_DECIMAL and isinstance(cell.value, float):
                    cell.number_format = "0.0"
                elif name in COUNTING and isinstance(cell.value, float):
                    cell.number_format = "0"
        note_sheet = writer.sheets["Σημείωση"]
        note_sheet.column_dimensions["A"].width = 36
        note_sheet.column_dimensions["B"].width = 110
        for cell in note_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
    print(f"wrote {OUTPUT} rows {len(frame)} missing {missing}", flush=True)


if __name__ == "__main__":
    build()
